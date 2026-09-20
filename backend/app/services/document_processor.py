"""
Learning Resource Processing (Module 10 in spec): extracts raw text from
uploaded PDFs / PPTs, and splits it into overlapping chunks suitable for
semantic embedding + retrieval (Module 3: AI Learning Repository).
"""
from typing import List
from pypdf import PdfReader
from pptx import Presentation


def extract_text_from_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(pages).strip()


def extract_text_from_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    slides_text = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        parts.append(text)
        slides_text.append("\n".join(parts))
    return "\n\n".join(slides_text).strip()


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> List[str]:
    """
    Simple sliding-window chunker over whitespace-split tokens. Overlap keeps
    context continuous across chunk boundaries, which improves retrieval
    quality for the RAG engine.
    """
    words = text.split()
    if not words:
        return []

    chunks = []
    start = 0
    while start < len(words):
        end = start + chunk_size
        chunk = " ".join(words[start:end])
        if chunk.strip():
            chunks.append(chunk.strip())
        if end >= len(words):
            break
        start = end - overlap
    return chunks


def process_file(file_path: str, material_type: str) -> str:
    if material_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif material_type == "ppt":
        return extract_text_from_pptx(file_path)
    else:
        raise ValueError(f"Unsupported material_type for text extraction: {material_type}")
